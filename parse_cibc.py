#!/usr/bin/env python3
"""Parse CIBC-DAILY-ETF-HOLDING-EN.xlsx (one sheet per ticker: row1 as-of, row2
headers 'Security Name'|'Weight %', rows3+ holdings) into cibc_holdings_raw.csv.
Sheets carry no holding tickers -> HOLDING_TICKER left blank."""
import csv, glob, os
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, 'cibc_holdings_raw.csv')

NAMES = {
    'CACB': 'CIBC Active Investment Grade Corporate Bond ETF',
    'CACE': 'Avantis CIBC Canadian Equity ETF',
    'CADE': 'Avantis CIBC International Equity ETF',
    'CAEM': 'Avantis CIBC Emerging Markets Equity ETF',
    'CAGE': 'Avantis CIBC All-Equity Asset Allocation ETF',
    'CAGR': 'Avantis CIBC Growth Asset Allocation ETF',
    'CAGX': 'Avantis CIBC World Equity ETF',
    'CAKE': 'Avantis CIBC Balanced Asset Allocation ETF',
    'CALB': 'CIBC Canadian Government Long-Term Bond ETF',
    'CALV': 'Avantis CIBC U.S. Large Cap Value ETF',
    'CASV': 'Avantis CIBC Global Small Cap Value ETF',
    'CAUS': 'Avantis CIBC U.S. All-Cap Equity ETF',
    'CAUV': 'Avantis CIBC U.S. Small Cap Value ETF',
    'CBLN': 'CIBC Balanced ETF Portfolio',
    'CCAD': 'CIBC Premium Cash Management ETF',
    'CCBI': 'CIBC Canadian Bond Index ETF',
    'CCCB': 'CIBC Canadian Banks Covered Call ETF',
    'CCDC': 'CIBC Canadian High Dividend Covered Call ETF',
    'CCEI': 'CIBC MSCI Canada Equity Index ETF',
    'CCLN': 'CIBC Clean Energy Index ETF',
    'CCLO': 'CIBC Income Advantage Fund - ETF Series',
    'CCNS': 'CIBC Conservative Fixed Income Pool - ETF Series',
    'CCON': 'CIBC Conservative ETF Portfolio',
    'CCRE': 'CIBC Core Fixed Income Pool - ETF Series',
    'CEMI': 'CIBC MSCI Emerging Markets Equity Index ETF',
    'CEQY': 'CIBC All-Equity ETF Portfolio',
    'CFRN': 'CIBC Active Investment Grade Floating Rate Bond ETF',
    'CGBI': 'CIBC Global Bond ex-Canada Index ETF (CAD-Hedged)',
    'CGRW': 'CIBC Balanced Growth ETF Portfolio',
    'CIEH': 'CIBC MSCI EAFE Equity Index ETF (CAD-Hedged)',
    'CIEI': 'CIBC MSCI EAFE Equity Index ETF',
    'CLBF': 'CIBC 1-5 Year Laddered Investment Grade Bond Fund - ETF Series',
    'CPLS': 'CIBC Core Plus Fixed Income Pool - ETF Series',
    'CQLC': 'CIBC Qx Canadian Low Volatility Dividend ETF',
    'CQLU': 'CIBC Qx U.S. Low Volatility Dividend ETF',
    'CSBA': 'CIBC Sustainable Balanced Solution - ETF Series',
    'CSBG': 'CIBC Sustainable Balanced Growth Solution - ETF Series',
    'CSBI': 'CIBC Canadian Short-Term Bond Index ETF',
    'CSCB': 'CIBC Sustainable Conservative Balanced Solution - ETF Series',
    'CSCE': 'CIBC Sustainable Canadian Equity Fund - ETF Series',
    'CSCP': 'CIBC Sustainable Canadian Core Plus Bond Fund - ETF Series',
    'CTBB': 'CIBC 2026 Investment Grade Bond Fund - ETF Series',
    'CTBC': 'CIBC 2027 Investment Grade Bond Fund - ETF Series',
    'CTBD': 'CIBC 2028 Investment Grade Bond Fund - ETF Series',
    'CTBE': 'CIBC 2029 Investment Grade Bond Fund - ETF Series',
    'CTBF': 'CIBC 2030 Investment Grade Bond Fund - ETF Series',
    'CTBG': 'CIBC 2031 Investment Grade Bond Fund - ETF Series',
    'CTUD.U': 'CIBC 2026 U.S. Investment Grade Bond Fund - ETF Series',
    'CTUE.U': 'CIBC 2027 U.S. Investment Grade Bond Fund - ETF Series',
    'CTUF.U': 'CIBC 2028 U.S. Investment Grade Bond Fund - ETF Series',
    'CTUG.U': 'CIBC 2029 U.S. Investment Grade Bond Fund - ETF Series',
    'CTUH.U': 'CIBC 2030 U.S. Investment Grade Bond Fund - ETF Series',
    'CTUI.U': 'CIBC 2031 U.S. Investment Grade Bond Fund - ETF Series',
    'CUDC': 'CIBC US High Dividend Covered Call ETF',
    'CUDC.F': 'CIBC US High Dividend Covered Call ETF (CAD-Hedged)',
    'CUEH': 'CIBC MSCI USA Equity Index ETF (CAD-Hedged)',
    'CUEI': 'CIBC MSCI USA Equity Index ETF',
    'CUSD.U': 'CIBC USD Premium Cash Management ETF',
}

def main():
    files = sorted(glob.glob(os.path.expanduser('~/workspace/browser_downloads/sess-*/browser-download-*-CIBC-DAILY-ETF-HOLDING-EN.xlsx')))
    path = files[-1]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        as_of = ws.cell(1, 2).value
        as_of = str(as_of)[:10] if as_of else ''
        n = 0
        for row in ws.iter_rows(min_row=3, values_only=True):
            name, wt = row[0], row[1]
            if name is None or wt is None:
                continue
            name = str(name).strip()
            try:
                w = float(wt)
            except (TypeError, ValueError):
                continue
            rows.append((sheet, NAMES.get(sheet, ''), '', name, '%.4f' % w, as_of))
            n += 1
        print(sheet, as_of, n, 'rows')
    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['TICKER', 'ETF_NAME', 'HOLDING_TICKER', 'HOLDING_NAME', 'WEIGHT_PCT', 'AS_OF_DATE'])
        w.writerows(rows)
    print('Wrote', OUT, len(rows), 'rows')

if __name__ == '__main__':
    main()
