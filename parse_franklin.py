#!/usr/bin/env python3
"""Parse Franklin PortfolioHoldingDetails.xlsx files (row8: fund name + as-of date,
row10: Security Name/Weight (%), rows11+ data) into franklin_holdings_raw.csv.
Files carry no holding tickers -> HOLDING_TICKER blank."""
import csv, glob, os
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, 'franklin_holdings_raw.csv')

FUND_TO_TICKER = {
    'Franklin All-Equity ETF Portfolio': 'EQY',
    'Franklin Brandywine Global Income Optimiser Fund': 'FBGO',
    'Franklin ClearBridge Global Infrastructure Income Fund': 'FCII',
    'Franklin Canadian Core Equity Fund': 'FCRC',
    'Franklin International Core Equity Fund': 'FCRI',
    'Franklin U.S. Core Equity Fund': 'FCRU',
    'Franklin U.S. Quality Moat Dividend Index ETF': 'FDIV',
    'Franklin Canadian Government Bond Fund': 'FGOV',
    'Franklin Canadian Ultra Short Term Bond Fund': 'FHIS',
    'Franklin FTSE India Index ETF': 'FID',
    'Franklin Innovation Fund': 'FINO',
    'Franklin FTSE U.S. Index ETF': 'FLAM',
    'Franklin FTSE Canada All Cap Index ETF': 'FLCD',
    'Franklin Canadian Corporate Bond Fund': 'FLCI',
    'Franklin Canadian Core Plus Bond Fund': 'FLCP',
    'Franklin Emerging Markets Equity Index ETF': 'FLEM',
    'Franklin Global Core Bond Fund': 'FLGA',
    'Franklin FTSE Japan Index ETF': 'FLJA',
    'Franklin Canadian Short Term Bond Fund': 'FLSD',
    'Franklin International Equity Index ETF': 'FLUR',
    'Franklin U.S. Large Cap Multifactor Index ETF': 'FLUS',
    'Franklin Canadian Low Volatility High Dividend Index ETF': 'FLVC',
    'Franklin International Low Volatility High Dividend Index ETF': 'FLVI',
    'Franklin U.S. Low Volatility High Dividend Index ETF': 'FLVU',
    'Franklin U.S. Mid Cap Multifactor Index ETF': 'FMID',
}

def norm_date(v):
    s = str(v).strip()
    if '/' in s:
        m, d, y = s.split('/')
        return '%04d-%02d-%02d' % (int(y), int(m), int(d))
    return s[:10]

def main():
    rows = []
    for path in sorted(glob.glob(os.path.expanduser('~/workspace/browser_downloads/sess-*/browser-download-*-PortfolioHoldingDetails.xlsx'))):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        fund = str(ws.cell(8, 1).value or '').strip()
        as_of = norm_date(ws.cell(8, 3).value)
        ticker = FUND_TO_TICKER.get(fund, '')
        n = 0
        for row in ws.iter_rows(min_row=11, values_only=True):
            name, wt = row[0], row[1]
            if name is None or wt is None:
                continue
            name = str(name).strip()
            if not name or 'IRC Fee' in name or 'Net Current Assets' in name:
                continue
            try:
                w = float(wt)
            except (TypeError, ValueError):
                continue
            rows.append((ticker, fund, '', name, '%.4f' % w, as_of))
            n += 1
        print(ticker or '?', fund[:45], as_of, n, 'rows')
    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['TICKER', 'ETF_NAME', 'HOLDING_TICKER', 'HOLDING_NAME', 'WEIGHT_PCT', 'AS_OF_DATE'])
        w.writerows(rows)
    print('Wrote', OUT, len(rows), 'rows')

if __name__ == '__main__':
    main()
